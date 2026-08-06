#!/usr/bin/env python3
"""One-shot import of the local Excel tracker into the hosted SQL tracker.

The Excel workbook (``Job_Application_Tracker.xlsx``) is the plugin/local
system of record.  The hosted product stores applications in SQL, so an owner
moving to the hosted app runs this once to carry their history across.

Safe to re-run: an application already present for the same
(user, company, title, applied date) is skipped rather than duplicated.

    python3 scripts/import_tracker_xlsx.py \
        --xlsx Job_Application_Tracker.xlsx --user-id 7 --dry-run
"""

from __future__ import annotations

import argparse
import sqlite3
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

# Column header -> position is resolved by name so a re-ordered or extended
# workbook still imports correctly.
COLUMNS = (
    "Company", "Job Title", "Application Date", "Status", "Resume File",
    "Cover Letter File", "Job Description", "ATS Score", "HR Score", "Notes",
    "Interview Date", "Follow Up Date", "Response", "Target Tier", "Fit Label",
    "Hard Reqs Missed", "Referral Source", "Rejection Reason",
    "Days To Response", "Interview Stages Reached",
)

# Mirrors the API-layer taxonomy (SQLite ALTER cannot enforce CHECK, so the
# importer validates too).  Anything else is preserved in notes instead of
# corrupting the values the insights aggregates depend on.
VALID_REJECTION_REASONS = frozenset({
    "no_response", "auto_reject", "screen_reject", "interview_reject",
    "offer_declined", "withdrawn",
})
VALID_TIERS = frozenset({"IC", "Sr", "Manager", "AD", "Director"})
VALID_FIT_LABELS = frozenset({"MEETS", "STRETCH", "MISS"})


def parse_score(value: Any) -> float | None:
    """Read a score cell that may be a number or a string like ``75.1%``."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().rstrip("%").strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def parse_stage(value: Any) -> int | None:
    """Interview stage on the documented 0-5 scale, or None."""
    if value is None or isinstance(value, bool):
        return None
    try:
        stage = int(str(value).strip())
    except (TypeError, ValueError):
        return None
    return stage if 0 <= stage <= 5 else None


def parse_count(value: Any) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        count = int(str(value).strip())
    except (TypeError, ValueError):
        return None
    return count if count >= 0 else None


def parse_date(value: Any) -> str | None:
    """Normalize a date cell to ISO ``YYYY-MM-DD``."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return text  # keep the raw value rather than silently dropping history


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _row_to_application(row: dict[str, Any]) -> dict[str, Any] | None:
    company = _clean(row.get("Company"))
    title = _clean(row.get("Job Title"))
    if not company or not title:
        return None

    notes_parts = [p for p in (_clean(row.get("Notes")),) if p]

    reason = _clean(row.get("Rejection Reason"))
    if reason and reason not in VALID_REJECTION_REASONS:
        notes_parts.append(f"Rejection reason (imported): {reason}")
        reason = None

    tier = _clean(row.get("Target Tier"))
    if tier and tier not in VALID_TIERS:
        notes_parts.append(f"Target tier (imported): {tier}")
        tier = None

    fit = _clean(row.get("Fit Label"))
    if fit:
        fit = fit.upper()
        if fit not in VALID_FIT_LABELS:
            notes_parts.append(f"Fit label (imported): {fit}")
            fit = None

    for label, key in (("Interview date", "Interview Date"),
                       ("Follow up", "Follow Up Date")):
        extra = parse_date(row.get(key))
        if extra:
            notes_parts.append(f"{label}: {extra}")

    return {
        "company": company,
        "job_title": title,
        "applied_at": parse_date(row.get("Application Date")),
        "status": _clean(row.get("Status")) or "Applied",
        "resume_file": _clean(row.get("Resume File")) or "",
        "cover_letter_file": _clean(row.get("Cover Letter File")) or "",
        "jd_ref": _clean(row.get("Job Description")),
        "ats_score": parse_score(row.get("ATS Score")),
        "hr_score": parse_score(row.get("HR Score")),
        "notes": "\n".join(notes_parts),
        "responded_at": parse_date(row.get("Response")),
        "target_tier": tier,
        "fit_label": fit,
        "hard_reqs_missed": parse_count(row.get("Hard Reqs Missed")),
        "referral_source": _clean(row.get("Referral Source")),
        "rejection_reason": reason,
        "interview_stage": parse_stage(row.get("Interview Stages Reached")),
    }


def _already_imported(conn: sqlite3.Connection, user_id: int, app: dict) -> bool:
    existing = conn.execute(
        "SELECT 1 FROM job_applications WHERE user_id = ? AND company = ? "
        "AND job_title = ? AND IFNULL(applied_at, '') = IFNULL(?, '') LIMIT 1",
        (user_id, app["company"], app["job_title"], app["applied_at"]),
    ).fetchone()
    return existing is not None


def import_workbook(
    xlsx_path: str | Path,
    *,
    user_id: int,
    conn: sqlite3.Connection,
    dry_run: bool = False,
) -> dict[str, int]:
    """Import applications for one user.  Returns counts by outcome."""
    import openpyxl  # imported lazily so the module loads without the dep

    workbook = openpyxl.load_workbook(str(xlsx_path), data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
    except StopIteration:
        return {"imported": 0, "skipped": 0, "blank": 0}

    stats = {"imported": 0, "skipped": 0, "blank": 0}
    for raw in rows:
        record = dict(zip(header, raw))
        app = _row_to_application(record)
        if app is None:
            stats["blank"] += 1
            continue
        if _already_imported(conn, user_id, app):
            stats["skipped"] += 1
            continue
        stats["imported"] += 1
        if dry_run:
            continue
        columns = ["user_id", *app.keys()]
        placeholders = ", ".join("?" for _ in columns)
        conn.execute(
            f"INSERT INTO job_applications ({', '.join(columns)}) "
            f"VALUES ({placeholders})",
            (user_id, *app.values()),
        )
    if not dry_run:
        conn.commit()
    return stats


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--user-id", required=True, type=int)
    parser.add_argument("--db", default=None, help="defaults to the app DB path")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)

    from cloud.db import get_conn, run_migrations

    db_path = args.db
    if db_path is None:
        from cloud.config import settings

        db_path = settings.DB_PATH
    conn = get_conn(db_path)
    run_migrations(conn)

    stats = import_workbook(
        args.xlsx, user_id=args.user_id, conn=conn, dry_run=args.dry_run
    )
    prefix = "Would import" if args.dry_run else "Imported"
    print(
        f"{prefix} {stats['imported']} application(s); "
        f"skipped {stats['skipped']} already present; "
        f"ignored {stats['blank']} blank row(s)."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
